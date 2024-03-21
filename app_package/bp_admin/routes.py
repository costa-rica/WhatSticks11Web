
from flask import Blueprint
from flask import render_template, url_for, redirect, flash, request, \
    abort, Response, current_app, send_from_directory, make_response
import bcrypt
from flask_login import login_required, login_user, logout_user, current_user
import logging
from logging.handlers import RotatingFileHandler
import os
import json
from ws_models import session_scope, engine, text, Users, Base
from app_package.bp_users.utils import userPermission

import pandas as pd
import shutil
from datetime import datetime
import openpyxl
import zipfile

from ws_utilities import create_df_crosswalk, update_and_append_user_location_day, \
    update_and_append_via_df_crosswalk_users, update_and_append_via_df_crosswalk_locations


#Setting up Logger
formatter = logging.Formatter('%(asctime)s:%(name)s:%(message)s')
formatter_terminal = logging.Formatter('%(asctime)s:%(filename)s:%(name)s:%(message)s')

#initialize a logger
logger_bp_admin = logging.getLogger(__name__)
logger_bp_admin.setLevel(logging.DEBUG)

file_handler = RotatingFileHandler(os.path.join(os.environ.get('WEB_ROOT'),'logs','bp_admin.log'), mode='a', maxBytes=5*1024*1024,backupCount=2)
file_handler.setFormatter(formatter)

#where the stream_handler will print
stream_handler = logging.StreamHandler()
stream_handler.setFormatter(formatter_terminal)

# logger_sched.handlers.clear() #<--- This was useful somewhere for duplicate logs
logger_bp_admin.addHandler(file_handler)
logger_bp_admin.addHandler(stream_handler)


salt = bcrypt.gensalt()


bp_admin = Blueprint('bp_admin', __name__)

@bp_admin.before_request
def before_request():
    logger_bp_admin.info(f"-- ***** in before_request route --")
    ###### TEMPORARILY_DOWN: redirects to under construction page ########
    if os.environ.get('TEMPORARILY_DOWN') == '1':
        if request.url != request.url_root + url_for('bp_main.temporarily_down')[1:]:
            # logger_bp_users.info("*** (logger_bp_users) Redirected ")
            logger_bp_admin.info(f'- request.referrer: {request.referrer}')
            logger_bp_admin.info(f'- request.url: {request.url}')
            return redirect(url_for('bp_main.temporarily_down'))

@bp_admin.route('/admin_page', methods = ['GET', 'POST'])
@login_required
def admin_page():
    with session_scope() as session:
        users_list=[i for i in session.query(Users).all()]
        users_user_loc_day_tuples = get_user_loc_day_tuple(users_list)
        test_flight_link = ""

        try:
            with open(os.path.join(current_app.config.get('WEBSITE_FILES'),'TestFlightUrl.txt'), 'r') as file:
                test_flight_link = file.read().strip()  # .strip() removes any leading/trailing whitespace
        except FileNotFoundError:
            logger_bp_admin.info(f"- no TestFlight Link found")

        if request.method == 'POST':
            formDict = request.form.to_dict()
            print('formDict:::', formDict)
            # if formDict.get('add_link'):
            if 'add_link' in formDict.keys():
                print("**** add ing link ****")
                if formDict.get('input_test_flight_link')[:8]=="https://":
                    with open(os.path.join(current_app.config.get('WEBSITE_FILES'),'TestFlightUrl.txt'), 'w') as file:
                        file.write(formDict.get('input_test_flight_link'))
                    flash(f'Updated TestFlight link', 'success')
                else:
                    flash(f'TestFlight link must be a valid https url', 'warning')    
                return redirect(url_for('bp_admin.admin_page'))

            # elif formDict.get('delete_link'):
            elif "delete_link" in formDict.keys():
                print("**** deleting link ****")
                os.remove(os.path.join(current_app.config.get('WEBSITE_FILES'),'TestFlightUrl.txt'))
                test_flight_link=""
                # flash(f'Removed TestFlight link from {request.url}', 'warning')
                flash(f'Removed TestFlight link from {request.host}', 'warning')


        return render_template('admin/admin.html', 
            # users_list=users_list, 
            users_user_loc_day_tuples=users_user_loc_day_tuples,
            test_flight_link=test_flight_link, str=str)


@bp_admin.route('/admin_db_download', methods = ['GET', 'POST'])
@login_required
def admin_db_download():
    logger_bp_admin.info('- in admin_db_download -')
    logger_bp_admin.info(f"current_user.admin_users_permission: {current_user.admin_users_permission}")

    if not current_user.admin_users_permission:
        return redirect(url_for('bp_main.home'))

    metadata = Base.metadata
    db_table_list = [table for table in metadata.tables.keys()]

    csv_dir_path = os.path.join(current_app.config.get('DATABASE_HELPER_FILES'), 'db_backup')

    if request.method == "POST":
        formDict = request.form.to_dict()

        if os.path.exists(os.path.join(current_app.config.get('DATABASE_HELPER_FILES'),"db_backup")):
            for filename in os.listdir(os.path.join(current_app.config.get('DATABASE_HELPER_FILES'),"db_backup")):
                os.remove(os.path.join(current_app.config.get('DATABASE_HELPER_FILES'),"db_backup",filename))

        # craete folder to save
        if not os.path.exists(os.path.join(current_app.config.get('DATABASE_HELPER_FILES'),"db_backup")):
            os.makedirs(os.path.join(current_app.config.get('DATABASE_HELPER_FILES'),"db_backup"))


        db_table_list = []
        for key, value in formDict.items():
            if value == "db_table":
                db_table_list.append(key)
      
        db_tables_dict = {}
        for table_name in db_table_list:
            with session_scope() as session:
                base_query = session.query(metadata.tables[table_name])
                df = pd.read_sql(text(str(base_query)), engine.connect())

            # fix table names
            cols = list(df.columns)
            for col in cols:
                if col[:len(table_name)] == table_name:
                    df = df.rename(columns=({col: col[len(table_name)+1:]}))


            db_tables_dict[table_name] = df
            if formDict.get("download_files") == "csv":
                db_tables_dict[table_name].to_csv(os.path.join(csv_dir_path, f"{table_name}.csv"), index=False)
            elif formDict.get("download_files") == "pickle":
                db_tables_dict[table_name].to_pickle(os.path.join(csv_dir_path, f"{table_name}.pkl"))
        
        shutil.make_archive(csv_dir_path, 'zip', csv_dir_path)

        return redirect(url_for('bp_admin.download_db_tables_zip'))
    
    return render_template('admin/admin_db_download_page.html', db_table_list=db_table_list )


@bp_admin.route('/admin_db_upload_single_file', methods = ['GET', 'POST'])
@login_required
def admin_db_upload_single_file():
    logger_bp_admin.info('- in admin_db_upload_single_file -')
    logger_bp_admin.info(f"current_user.admin: {current_user.admin_users_permission}")

    if not current_user.admin_users_permission:
        return redirect(url_for('bp_main.home'))

    metadata = Base.metadata
    db_table_list = [table for table in metadata.tables.keys()]
    
    list_files_in_db_upload = os.listdir(current_app.config.get('DB_UPLOAD'))
    list_files_in_db_upload_csv_pkl_zip = []
    for filename_string in list_files_in_db_upload:
        filename, file_extension = os.path.splitext(filename_string)
        if file_extension in ['.csv','.pkl']:
            list_files_in_db_upload_csv_pkl_zip.append(filename_string)

    if request.method == "POST":
        formDict = request.form.to_dict()
        print(f"- admin_db_upload_single_file POST -")
        print("formDict: ", formDict)


        if formDict.get('what_kind_of_post') == "upload_from_here":
            requestFiles = request.files

            
            # csv_file_for_table = request.files.get('csv_table_upload')
            file_for_table_upload = request.files.get('file_for_table_upload')
            file_for_table_upload_filename = file_for_table_upload.filename

            logger_bp_admin.info(f"-- Get CSV file name --")
            logger_bp_admin.info(f"--  {file_for_table_upload_filename} --")

            ## save to databases/WhatSticks/database_helpers/DB_UPLOAD directory
            path_to_uploaded_table_file = os.path.join(current_app.config.get('DB_UPLOAD'),file_for_table_upload_filename)
            file_for_table_upload.save(path_to_uploaded_table_file)

            logger_bp_admin.info(f"-- table to go to: { formDict.get('existing_db_table_to_update')}")

            return redirect(url_for('bp_admin.upload_table', table_name = formDict.get('existing_db_table_to_update'),
                path_to_uploaded_table_file=path_to_uploaded_table_file))
                # path_to_uploaded_csv=path_to_uploaded_csv))
        elif formDict.get('what_kind_of_post') == "uploaded_already":
            already_uploaded_filename = formDict.get('selectedFile')
            path_to_uploaded_table_file = os.path.join(current_app.config.get('DB_UPLOAD'),already_uploaded_filename)
            
            return redirect(url_for('bp_admin.upload_table', table_name = formDict.get('existing_db_table_to_update'),
                path_to_uploaded_table_file=path_to_uploaded_table_file))

    return render_template('admin/admin_db_upload_single_file_page.html', db_table_list=db_table_list,
        len=len, list_files_in_db_upload_csv_pkl_zip=list_files_in_db_upload_csv_pkl_zip)


@bp_admin.route('/upload_table/<table_name>', methods = ['GET', 'POST'])
@login_required
def upload_table(table_name):
    logger_bp_admin.info('- in upload_table -')
    logger_bp_admin.info(f"current_user.admin: {current_user.admin_users_permission}")
    # path_to_uploaded_csv = request.args.get('path_to_uploaded_csv')
    path_to_uploaded_table_file = request.args.get('path_to_uploaded_table_file')

    if not current_user.admin_users_permission:
        return redirect(url_for('bp_main.home'))

    # Get Table Column names from the database corresponding to the running webiste/app
    metadata = Base.metadata
    existing_table_column_names = metadata.tables[table_name].columns.keys()

    if path_to_uploaded_table_file[-4:]=='.pkl':
        df = pd.read_pickle(path_to_uploaded_table_file)
    else:
        # Get column names from the uploaded csv
        df = pd.read_csv(path_to_uploaded_table_file)

    if 'time_stamp_utc' in df.columns and path_to_uploaded_table_file[-4:]!='.pkl':
        try:
            df['time_stamp_utc'] = pd.to_datetime(df['time_stamp_utc'], format='%d/%m/%Y %H:%M')
        # except ValueError:
        #     df['time_stamp_utc'] = pd.to_datetime(df['time_stamp_utc'], format='%d/%m/%Y %H:%M:%S')
        except:
            # df = pd.read_csv(path_to_uploaded_csv, parse_dates=['time_stamp_utc'])
            df = pd.read_csv(path_to_uploaded_table_file, parse_dates=['time_stamp_utc'])
    replacement_data_col_names = list(df.columns)
    # Match column names between the two tables
    match_cols_dict = {}
    for existing_db_column in existing_table_column_names:
        try:
            index = replacement_data_col_names.index(existing_db_column)
            match_cols_dict[existing_db_column] = replacement_data_col_names[index]
        except ValueError:
            match_cols_dict[existing_db_column] = None

    if request.method == "POST":
        formDict = request.form.to_dict()


        # NOTE: upload data to existing database
        ### formDict (key) is existing databaes column name
        # existing_names_list = [existing for existing, update in formDict.items() if update != 'true' ]
        
        # check for default values and remove from formDict
        set_default_value_dict = {}
        for key, value in formDict.items():
            if key[:len("default_checkbox_")] == "default_checkbox_":
                set_default_value_dict[value] = formDict.get(value)
        

        # Delete elements from dictionary
        for key, value in set_default_value_dict.items():
            del formDict[key]
            checkbox_key = "default_checkbox_" + key
            del formDict[checkbox_key]

        existing_names_list = []
        for key, value in formDict.items():
            if value != 'true':
                existing_names_list.append(key)
            
        df_update = pd.DataFrame(columns=existing_names_list)

        # value is the new data (aka the uploaded csv file column)
        for exisiting, replacement in formDict.items():
            if not replacement in ['true','']:
                # print(replacement)
                df_update[exisiting]=df[replacement].values

        # Add in columns with default values
        for column_name, default_value in set_default_value_dict.items():
            if column_name == 'time_stamp_utc': 
                df_update[column_name] = datetime.utcnow()
            else:
                df_update[column_name] = default_value

        
        # remove existing users from upload
        # NOTE: There needs to be a user to upload data
        if table_name == 'users':
            print("--- Found users table ---")
            with session_scope() as session:
                existing_users = session.query(Users).all()
                list_of_emails_in_db = [i.email for i in existing_users]
                for email in list_of_emails_in_db:
                    # df_update.drop(df_update[df_update.email== email].index, inplace = True)
                    print(f"-- removing {email} from upload dataset --")
                    df_update.drop(df_update[df_update.email == email].index, inplace=True)
                    df_update.reset_index(drop=True, inplace=True)

            # Assuming 'password' column exists and you want to encode all non-null passwords
            if 'password' in df_update.columns:
                df_update['password'] = df_update['password'].apply(lambda x: x.encode() if pd.notnull(x) else x)

        df_update.to_sql(table_name, con=engine, if_exists='append', index=False)

        flash(f"{table_name} update: successful!", "success")

        # return redirect(request.url)
        return redirect(url_for('bp_admin.admin_db_upload_single_file'))
    
    return render_template('admin/upload_table.html', table_name=table_name, 
        match_cols_dict = match_cols_dict,
        existing_table_column_names=existing_table_column_names,
        replacement_data_col_names = replacement_data_col_names)


@bp_admin.route('/admin_db_upload_zip', methods = ['GET', 'POST'])
@login_required
def admin_db_upload_zip():
    logger_bp_admin.info('- in admin_db_upload_zip -')

    if not current_user.admin_users_permission:
        return redirect(url_for('bp_main.home'))

    metadata = Base.metadata
    db_table_list = [table for table in metadata.tables.keys()]
    
    list_files_in_db_upload = os.listdir(current_app.config.get('DB_UPLOAD'))
    list_files_in_db_upload_csv_pkl_zip = []
    for filename_string in list_files_in_db_upload:
        filename, file_extension = os.path.splitext(filename_string)

        if file_extension in ['.zip']:
            list_files_in_db_upload_csv_pkl_zip.append(filename_string)

    if request.method == "POST":
        formDict = request.form.to_dict()
        logger_bp_admin.info(f"- admin_db_upload_zip POST -")
        logger_bp_admin.info(f"formDict: {formDict}")
        logger_bp_admin.info(f"request.files: {request.files}")
        
        file_for_table_upload = request.files.get('zip_filename_uploaded')
        


        if file_for_table_upload.filename != '':

            zip_filename = file_for_table_upload.filename

            logger_bp_admin.info(f"-- Get .zip file name --")
            logger_bp_admin.info(f"-- {zip_filename} --")

            ## save to databases/WhatSticks/database_helpers/DB_UPLOAD directory
            path_to_uploaded_table_file = os.path.join(current_app.config.get('DB_UPLOAD'),zip_filename)
            file_for_table_upload.save(path_to_uploaded_table_file)
        else:
            zip_filename = formDict.get('zip_filename_existing')

        if zip_filename == None:
            flash(f"Select a .zip file", "warning")
            return redirect(request.referrer)
        count_of_new_users = 0
        count_of_new_locations = 0
        count_of_new_user_location_day = 0
        count_of_new_weather_hist = 0
        count_of_new_workouts = 0
        count_of_new_qty_cat = 0

        # Step 1: Make Crosswalks 
        df_crosswalk_users = create_df_crosswalk('users', zip_filename)
        df_crosswalk_locations = create_df_crosswalk('locations', zip_filename)

        logger_bp_admin.info(f"-- count of df_crosswalk_users {len(df_crosswalk_users)} --")
        logger_bp_admin.info(f"-- count of df_crosswalk_locations {len(df_crosswalk_locations)} --")

        if 'new_row' in df_crosswalk_users.columns:
            count_of_new_users = len(df_crosswalk_users[df_crosswalk_users.new_row=='yes'])
            df_crosswalk_users.drop(columns=['new_row'], inplace=True)
        if 'new_row' in df_crosswalk_locations.columns:
            count_of_new_locations = len(df_crosswalk_locations[df_crosswalk_locations.new_row=='yes'])
            df_crosswalk_locations.drop(columns=['new_row'], inplace=True)

        if len(df_crosswalk_locations) > 0 and len(df_crosswalk_users) > 0:
            # Step 2: Add UserLocationDay data with new user_ids and location_ids, if any new rows to add
            count_of_new_user_location_day = update_and_append_user_location_day(
                zip_filename,df_crosswalk_users,df_crosswalk_locations)
        
        if len(df_crosswalk_locations) > 0:
            # Step 3: Add WeatherHistory with new location_ids, if any new rows to add
            count_of_new_weather_hist = update_and_append_via_df_crosswalk_locations(
                'weather_history', 'location_id', zip_filename,df_crosswalk_locations)

        if len(df_crosswalk_users) > 0:
            # Step 4: Add AppleHealthWorkouts with new user_ids, if any new rows to add
            count_of_new_workouts = update_and_append_via_df_crosswalk_users(
                'apple_health_workout',zip_filename,df_crosswalk_users)
            
            # Step 5: Add AppleHealthQuantityCategory with new user_ids, if any new rows to add
            count_of_new_qty_cat = update_and_append_via_df_crosswalk_users(
                'apple_health_quantity_category',zip_filename,df_crosswalk_users)

        # request.referrer - the url for the page that sent 
        # in this case it's just this same page.

        long_f_string = (
            f"Successfully added: " +
            f"\n Users....................... {count_of_new_users:,} " +
            f"\n Locations................... {count_of_new_locations:,}" +
            f"\n UserLocationDay............. {count_of_new_user_location_day:,}" +
            f"\n WeatherHistory.............. {count_of_new_weather_hist:,}" +
            f"\n Workouts.................... {count_of_new_workouts:,}" +
            f"\n AppleHealthQuantityCategory. {count_of_new_qty_cat:,}"
        )

        flash( long_f_string, "success")
        return redirect(request.referrer)

    return render_template('admin/admin_db_upload_zip_page.html', db_table_list=db_table_list,
        len=len, list_files_in_db_upload_csv_pkl_zip=list_files_in_db_upload_csv_pkl_zip)



@bp_admin.route('/nrodrig1_admin', methods=["GET"])
def nrodrig1_admin():
    with session_scope() as session:
        nrodrig1 = session.query(Users).filter_by(email="nrodrig1@gmail.com").first()
        if nrodrig1 != None:
            nrodrig1.admin_users_permission = True
            flash("nrodrig1@gmail updated to admin", "success")
        return redirect(url_for('bp_main.home'))


# @bp_admin.route("/download_db_tables_as_csv", methods=["GET","POST"])
@bp_admin.route("/download_db_tables_zip", methods=["GET","POST"])
@login_required
def download_db_tables_zip():
    return send_from_directory(os.path.join(current_app.config['DATABASE_HELPER_FILES']),'db_backup.zip', as_attachment=True)



@bp_admin.route("/delete_db_upload_file/<filename>", methods=["GET","POST"])
@login_required
def delete_db_upload_file(filename):
    logger_bp_admin.info(f"- accessed delete_db_upload_file -")
    logger_bp_admin.info(f"- deleting {filename} -")

    try:
        os.remove(os.path.join(current_app.config['DB_UPLOAD'],filename))
    except FileNotFoundError:
        logger_bp_admin.info(f"File {filename} was not found.")
    except PermissionError:
        logger_bp_admin.info(f"Permission denied: Unable to delete {filename}.")
    except Exception as e:
        logger_bp_admin.info(f"An error occurred: {e}")
    

    # Redirect back to the referrer page
    referrer = request.referrer
    if referrer:
        return redirect(referrer)
    else:
        # Fallback to a default route if the referrer is not available
        return redirect(url_for('bp_admin.admin_page'))



@bp_admin.route("/delete_user/<email>", methods=["GET","POST"])
@login_required
def delete_user(email):
    logger_bp_admin.info(f'-accessed: delete_user {email}')
    # with open(os.path.join(current_app.config['DIR_DB_FILES_UTILITY'],'added_users.txt')) as json_file:
    #     get_users_dict=json.load(json_file)
    #     json_file.close()
    
    # del get_users_dict[email]
    
    # added_users_file=os.path.join(current_app.config['DIR_DB_FILES_UTILITY'], 'added_users.txt')
    # with open(added_users_file, 'w') as json_file:
    #     json.dump(get_users_dict, json_file)
        
    # if len(sess.query(User).filter_by(email=email).all())>0:
    #     sess.query(User).filter_by(email=email).delete()
    #     sess_users.commit()
    
    
    
    # flash(f'{email} has been deleted!', 'success')
    flash(f'NOT deleteing emails with this link yet', 'success')
    return redirect(url_for('bp_admin.admin_page'))
    # return redirect(request.url)



